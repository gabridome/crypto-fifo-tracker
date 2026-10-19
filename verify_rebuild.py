#!/usr/bin/env python3
"""
Prove that two databases say the same thing — or show exactly where they differ.

Written for one question: *can the database be rebuilt from the exports in
`data/` alone?* `import_all_data.py` does the rebuild, this script proves the
result. It compares what matters and ignores only what cannot be compared, for
a stated reason:

  * **ignored**: `id` (AUTOINCREMENT, depends on insertion order), `created_at`
    and `imported_at` (wall clock of the import), and the foreign keys that
    carry ids.
  * **explained**: `fee_currency` empty in the reference and `'EUR'` in the
    rebuild. The live database predates the current importers on most rows;
    `doc/schema.sql` declares `fee_currency TEXT DEFAULT 'EUR'`, so a rebuild
    fills in what those old rows left NULL. The deroga is narrow on purpose —
    empty → `EUR` only. `USD` → `EUR` would be a conversion bug, and this
    script must not be the place where it goes unnoticed.
  * **compared**, row by row as a multiset, on the columns the two databases
    share: dates, types, amounts, prices, values, fees, currencies, exchanges,
    sources, hashes, swap groups, notes; every FIFO lot's date, quantity and
    cost; every sale match's amounts, cost basis, proceeds, gain and holding
    period; every deferred-swap lot transfer; the tax aggregates per year; the
    open lots per cryptocurrency. Columns present on one side only are reported
    separately and fail the check when they carry data — they are not silently
    dropped from the comparison.

With `--annex YEAR` it also generates the IRS annex from both databases and
compares it cell by cell (ignoring the generation timestamp), which closes the
chain the figures actually travel: exports → FIFO → annex. Both databases are
copied to a temporary directory first, so generating the annex never touches
`data/reports/`.

Usage:
    python3 verify_rebuild.py <riferimento.db> <ricostruito.db>
    python3 verify_rebuild.py <riferimento.db> <ricostruito.db> --annex 2024
    python3 verify_rebuild.py <un.db> --fingerprint      # impronta di UN database

Exit code 0 only when nothing unexplained differs.
"""

import hashlib
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from collections import Counter, namedtuple

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))

Check = namedtuple('Check', 'name status detail')
Report = namedtuple('Report', 'ok checks notes yearly')

# Colonne che non si possono confrontare, e perché.
UNCOMPARABLE = {
    'transactions': ('id', 'created_at', 'imported_at'),
    'fifo_lots': ('id', 'purchase_transaction_id', 'created_at'),
    'sale_lot_matches': ('id', 'sale_transaction_id', 'fifo_lot_id', 'created_at'),
    'swap_lot_transfers': ('id', 'source_fifo_lot_id', 'created_at'),
}
# Trattata a parte, con la sua regola stretta.
SPECIAL = {'transactions': ('fee_currency',)}

# Le colonne dell'impronta: elenco FISSO e ordinato, non «quelle che il database
# ha». Due ragioni, entrambe misurate il 2026-10-19: l'ordine delle colonne nel
# DDL del database vivo non è quello di doc/schema.sql, e un hash costruito
# sull'ordine di PRAGMA cambia a dati identici; e il database vivo porta tre
# colonne legacy vuote che un ricostruito non ha. Con un elenco fisso l'impronta
# del vivo e quella del ricostruito coincidono — che è l'unica proprietà per cui
# l'impronta serve.
FINGERPRINT_COLUMNS = {
    'transactions': sorted((
        'transaction_date', 'transaction_type', 'cryptocurrency', 'amount',
        'price_per_unit', 'total_value', 'fee_amount', 'currency',
        'exchange_name', 'source', 'record_hash', 'swap_group_id',
        'transaction_id', 'notes')),
    'fifo_lots': sorted((
        'cryptocurrency', 'purchase_date', 'original_amount', 'remaining_amount',
        'purchase_price_per_unit', 'cost_basis', 'purchase_fee_total',
        'exchange_name')),
    'sale_lot_matches': sorted((
        'sale_date', 'purchase_date', 'cryptocurrency', 'amount_sold',
        'purchase_price_per_unit', 'sale_price_per_unit', 'cost_basis',
        'proceeds', 'gain_loss', 'holding_period_days')),
    # Entrata nell'impronta il 2026-10-19: prima era fuori, e una review ha
    # mostrato che due database con `swap_lot_transfers` diverso in **ogni**
    # colonna davano la stessa impronta e lo stesso «✓ IDENTICI». Questa tabella
    # porta data d'acquisto e costo che attraversano una permuta differita,
    # cioe' proprio cio' che decide l'esenzione oltre i 365 giorni.
    'swap_lot_transfers': sorted((
        'swap_group_id', 'from_cryptocurrency', 'to_cryptocurrency', 'swap_date',
        'source_purchase_date', 'amount_from', 'cost_basis')),
}


def _ro(path):
    conn = sqlite3.connect(f'file:{path}?mode=ro', uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def _columns(conn, table):
    return [r[1] for r in conn.execute(f'PRAGMA table_info({table})')]


def _count(conn, table):
    has = conn.execute(
        "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name=?",
        (table,)).fetchone()[0]
    return conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] if has else None


def _multiset_hash(conn, table, cols):
    if not cols:
        raise ValueError(f'{table}: nessuna colonna da confrontare')
    quoted = ', '.join(f'"{c}"' for c in cols)
    rows = sorted(repr(tuple(r)) for r in conn.execute(f'SELECT {quoted} FROM "{table}"'))
    return hashlib.sha256('\n'.join(rows).encode()).hexdigest()


def _yearly(conn):
    q = """SELECT substr(sale_date,1,4) AS year,
                  COUNT(*)                                                   AS matches,
                  SUM(CASE WHEN holding_period_days>=365 THEN 1 ELSE 0 END)  AS exempt_n,
                  ROUND(SUM(CASE WHEN holding_period_days>=365
                                 THEN gain_loss ELSE 0 END),2)               AS exempt_gain,
                  ROUND(SUM(CASE WHEN holding_period_days<365
                                 THEN gain_loss ELSE 0 END),2)               AS taxable_gain
           FROM sale_lot_matches GROUP BY 1 ORDER BY 1"""
    return [dict(r) for r in conn.execute(q)]


def _open_lots(conn):
    q = """SELECT cryptocurrency,
                  COUNT(*)                                                   AS lots,
                  ROUND(SUM(remaining_amount),8)                             AS amount,
                  ROUND(SUM(cost_basis*remaining_amount/original_amount),2)   AS cost_eur
           FROM fifo_lots WHERE remaining_amount > 0 GROUP BY 1 ORDER BY 1"""
    return [dict(r) for r in conn.execute(q)]


def _fee_currency_check(a, b, cols):
    """Vuoto → EUR è la deroga. Qualunque altra differenza è un errore.

    Identificata **per riga**, non per distribuzione. La prima versione di
    questa funzione confrontava i conteggi `(source, fee_currency)` e il
    2026-10-19 una review l'ha bucata in un minuto: con una fonte che ha sia
    righe legacy vuote sia righe non-EUR, i conteggi tornano anche se una riga
    è passata da `USD` a `EUR` — cioè esattamente la firma dell'errore di
    conversione che questa funzione dichiarava di bocciare, riportata come
    «spiegato». Il docstring di allora sosteneva che due righe della stessa
    fonte «non si distinguono»: era falso, hanno date e hash diversi.

    Ora il criterio è letterale: prendi le righe del riferimento, sostituisci
    `''` con `'EUR'` nella sola colonna `fee_currency`, e i due multiset devono
    diventare **uguali**. Se non lo diventano, c'è una differenza che la deroga
    non copre, e si boccia mostrandola.
    """
    quoted = ', '.join(f'"{c}"' for c in cols)
    q = (f'SELECT {quoted}, COALESCE(fee_currency, \'\') AS fc '
         'FROM transactions')
    ca = Counter(tuple(r) for r in a.execute(q))
    cb = Counter(tuple(r) for r in b.execute(q))
    if ca == cb:
        return Check('fee_currency', 'identico', 'stesse valute delle commissioni')

    # La deroga: nel riferimento erano vuote, nel ricostruito sono EUR.
    moved = 0
    ca_filled = Counter()
    for row, n in ca.items():
        if row[-1] == '':
            moved += n
            row = row[:-1] + ('EUR',)
        ca_filled[row] += n

    if ca_filled == cb:
        return Check('fee_currency', 'spiegato',
                     f'{moved:,} righe con fee_currency vuota nel riferimento, '
                     "riempite a 'EUR' dal DEFAULT di doc/schema.sql — ogni altra "
                     'riga corrisponde come multiset. Nota: `COALESCE` confonde '
                     "NULL e stringa vuota, quindi un NULL→'' su questa colonna "
                     'è invisibile sia qui sia nel confronto principale, che la '
                     'esclude.')

    only_ref = ca_filled - cb
    only_new = cb - ca_filled
    sample = []
    for label, counter in (('solo nel riferimento', only_ref), ('solo nel ricostruito', only_new)):
        for row, n in list(counter.items())[:3]:
            sample.append(f'{label} (×{n}): fee_currency={row[-1]!r} su {row[:3]}')
    return Check('fee_currency', 'DIVERSO',
                 f'{sum(only_ref.values())} righe del riferimento e '
                 f'{sum(only_new.values())} del ricostruito non si corrispondono '
                 'nemmeno applicando la deroga vuoto→EUR:\n      '
                 + '\n      '.join(sample))


def _column_sets_check(a, b, tables):
    """Colonne presenti in un database e non nell'altro, in **entrambe** le
    direzioni e su **tutte** le tabelle confrontate.

    Due correzioni del 2026-10-19, trovate da una review. La prima versione
    guardava solo `transactions`: una colonna popolata persa in `fifo_lots` o in
    `sale_lot_matches` passava in silenzio, che è precisamente il difetto che
    questa funzione esiste per intercettare. E guardava in una sola direzione:
    una colonna presente solo nel ricostruito non veniva né confrontata né
    riportata, quindi ogni colonna aggiunta a `doc/schema.sql` nasceva esente
    da verifica.

    Regola: una colonna che manca da un lato e **contiene dati** dall'altro è
    una differenza di contenuto → boccia. Vuota → si riporta e si spiega.
    """
    notes, populated, empty = [], [], []
    for table in tables:
        ca, cb = _columns(a, table), _columns(b, table)
        for col, where, conn in ([(c, 'riferimento', a) for c in ca if c not in cb]
                                 + [(c, 'ricostruito', b) for c in cb if c not in ca]):
            n = conn.execute(
                f'SELECT COUNT(*) FROM "{table}" WHERE "{col}" IS NOT NULL').fetchone()[0]
            if n:
                distinct = [r[0] for r in conn.execute(
                    f'SELECT DISTINCT "{col}" FROM "{table}" '
                    f'WHERE "{col}" IS NOT NULL LIMIT 3')]
                populated.append((table, col, where, n))
                notes.append(f'colonna `{table}.{col}`: solo nel {where}, '
                             f'{n:,} valori (es. {distinct})')
            else:
                empty.append((table, col, where))
                notes.append(f'colonna legacy `{table}.{col}`: solo nel {where}, '
                             'interamente vuota — nessun dato perso')

    if populated:
        return Check('colonne presenti da un solo lato', 'DIVERSO',
                     '; '.join(f'{t}.{c} ({n:,} valori, solo nel {w})'
                               for t, c, w, n in populated)), notes
    if empty:
        return Check('colonne presenti da un solo lato', 'spiegato',
                     ', '.join(f'{t}.{c}' for t, c, _ in empty) + ' (vuote)'), notes
    return Check('colonne presenti da un solo lato', 'identico', 'nessuna'), notes


def compare(ref_db, new_db):
    """Confronta due database. Ritorna un `Report`; non stampa nulla."""
    a, b = _ro(ref_db), _ro(new_db)
    try:
        checks, notes = [], []

        for table in ('transactions', 'fifo_lots', 'sale_lot_matches',
                      'swap_lot_transfers'):
            na, nb = _count(a, table), _count(b, table)
            checks.append(Check(
                f'conteggio {table}',
                'identico' if na == nb else 'DIVERSO',
                f'riferimento {na:,} / ricostruito {nb:,}' if na is not None
                and nb is not None else f'riferimento {na} / ricostruito {nb}'))

        for table, label in (
                ('transactions', 'transazioni (colonne confrontabili)'),
                ('fifo_lots', 'lotti FIFO (date, quantità, costi)'),
                ('sale_lot_matches', 'match di vendita (importi, costi, gain, giorni)'),
                ('swap_lot_transfers', 'trasferimenti di lotto (permute differite)')):
            shared = [c for c in _columns(a, table) if c in _columns(b, table)
                      and c not in UNCOMPARABLE.get(table, ())
                      and c not in SPECIAL.get(table, ())]
            if not shared:
                # Tabella assente da un lato: il conteggio sopra l'ha già
                # bocciata, e questo confronto non ha nulla su cui lavorare.
                # Senza questo ramo si costruiva `SELECT  FROM "<tabella>"` e
                # l'intero strumento moriva con un OperationalError — perdendo
                # la diagnosi di tutto il resto proprio nel caso per cui
                # esiste: un database vecchio, da `data/backups`, che non ha
                # ancora `swap_lot_transfers`.
                checks.append(Check(
                    label, 'DIVERSO',
                    'tabella assente da un lato: nessuna colonna in comune, '
                    'contenuto non confrontabile'))
                continue
            ha, hb = _multiset_hash(a, table, shared), _multiset_hash(b, table, shared)
            checks.append(Check(
                label, 'identico' if ha == hb else 'DIVERSO',
                f'{len(shared)} colonne, hash {ha[:16]} vs {hb[:16]}'))

        fc_cols = [c for c in _columns(a, 'transactions')
                   if c in _columns(b, 'transactions')
                   and c not in UNCOMPARABLE['transactions']
                   and c not in SPECIAL['transactions']]
        checks.append(_fee_currency_check(a, b, fc_cols))
        cols_check, cols_notes = _column_sets_check(
            a, b, ('transactions', 'fifo_lots', 'sale_lot_matches',
                   'swap_lot_transfers'))
        checks.append(cols_check)
        notes.extend(cols_notes)
        notes.append('fee_currency: ' + [c for c in checks
                                         if c.name == 'fee_currency'][0].detail)

        ya, yb = _yearly(a), _yearly(b)
        checks.append(Check('aggregati fiscali per anno',
                            'identico' if ya == yb else 'DIVERSO',
                            f'{len(ya)} anni nel riferimento, {len(yb)} nel ricostruito'))

        oa, ob = _open_lots(a), _open_lots(b)
        checks.append(Check(
            'giacenze aperte',
            'identico' if oa == ob else 'DIVERSO',
            '; '.join(f'{r["cryptocurrency"]} {r["lots"]} lotti, {r["amount"]}, '
                      f'EUR {r["cost_eur"]:,.2f}' for r in oa)))

        ok = all(c.status != 'DIVERSO' for c in checks)
        return Report(ok=ok, checks=checks, notes=notes, yearly=ya)
    finally:
        a.close()
        b.close()


def fingerprint(db_path):
    """L'impronta di determinismo di UN database, da conservare col backup.

    `compare()` ha bisogno di due database. Dopo un disastro ce n'è uno solo —
    quello appena ricostruito — e non c'è niente con cui confrontarlo. Questa
    funzione produce le righe da archiviare **insieme** agli export: chi
    ricostruisce rilancia `--fingerprint` e verifica di aver ottenuto gli stessi
    numeri, senza bisogno del database di partenza.
    """
    a = _ro(db_path)
    try:
        lines = [f'database   {os.path.basename(db_path)}']
        for table in ('transactions', 'fifo_lots', 'sale_lot_matches',
                      'swap_lot_transfers'):
            n = _count(a, table)
            if n is None:
                raise RuntimeError(
                    f'tabella {table} assente dal database: l\'impronta non è '
                    'confrontabile con quella di un database completo, e non va '
                    'archiviata come se lo fosse.')
            lines.append(f'righe {table:20} {n}')
        for table, cols in FINGERPRINT_COLUMNS.items():
            missing = [c for c in cols if c not in _columns(a, table)]
            if missing:
                raise RuntimeError(
                    f'{table}: colonne dell\'impronta assenti dal database: '
                    f'{missing}. L\'impronta non è confrontabile e non va '
                    'archiviata come se lo fosse.')
            lines.append(f'hash  {table:20} {_multiset_hash(a, table, cols)[:32]}'
                         f'  ({len(cols)} colonne canoniche)')
        lines.append('')
        lines.append(f'{"anno":6} {"match":>7} {"esenti":>7} {"plus esente":>18} {"tassabile":>14}')
        for r in _yearly(a):
            lines.append(f'{r["year"]:6} {r["matches"]:>7} {r["exempt_n"]:>7} '
                         f'{r["exempt_gain"]:>18,.2f} {r["taxable_gain"]:>14,.2f}')
        lines.append('')
        for r in _open_lots(a):
            lines.append(f'aperti {r["cryptocurrency"]:6} {r["lots"]:>4} lotti  '
                         f'{r["amount"]:>18}  EUR {r["cost_eur"]:>16,.2f}')
        return lines
    finally:
        a.close()


# ── allegato IRS ─────────────────────────────────────────────────────────────

def _generate_annex(db_path, year, workdir):
    """Genera l'allegato da `db_path` in una copia temporanea del database.

    Il generatore scrive in `<cartella del db>/reports/`: senza la copia,
    verificare sovrascriverebbe i report veri in `data/reports/`.
    """
    sandbox = os.path.join(workdir, os.path.basename(db_path).replace('.db', ''))
    os.makedirs(sandbox, exist_ok=True)
    copy = os.path.join(sandbox, 'db.db')
    shutil.copy2(db_path, copy)
    result = subprocess.run(
        [sys.executable, os.path.join(PROJECT_ROOT, 'calculators/generate_irs_report.py'),
         str(year), copy],
        capture_output=True, text=True, cwd=PROJECT_ROOT)
    out = os.path.join(sandbox, 'reports', f'IRS_Crypto_FIFO_{year}.xlsx')
    if result.returncode != 0 or not os.path.exists(out):
        raise RuntimeError(
            f'generazione allegato {year} fallita per {db_path}:\n'
            + (result.stdout + result.stderr)[-800:])
    return out


def compare_annex(ref_db, new_db, year):
    """Confronta cella per cella i due allegati. Ritorna (ok, righe di esito)."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return False, ['openpyxl non installato: impossibile confrontare l\'allegato']

    lines = []
    with tempfile.TemporaryDirectory() as tmp:
        try:
            fa = _generate_annex(ref_db, year, tmp)
            fb = _generate_annex(new_db, year, tmp)
        except RuntimeError as e:
            return False, [str(e)]

        wa, wb = load_workbook(fa, data_only=True), load_workbook(fb, data_only=True)
        if wa.sheetnames != wb.sheetnames:
            return False, [f'fogli diversi: {wa.sheetnames} vs {wb.sheetnames}']

        diffs, cells, skipped = [], 0, 0
        for name in wa.sheetnames:
            sa, sb = wa[name], wb[name]
            rows_a = list(sa.iter_rows(values_only=True))
            rows_b = list(sb.iter_rows(values_only=True))
            if len(rows_a) != len(rows_b):
                diffs.append(f'{name}: {len(rows_a)} righe vs {len(rows_b)}')
                continue
            # strict=False di proposito su entrambi gli zip, con il controllo
            # delle lunghezze fatto a mano: una riga con un numero di colonne
            # diverso va **riportata** come differenza, non sollevata come
            # eccezione — e zip() senza controllo la troncherebbe in silenzio,
            # che è il modo di non vederla mai.
            for i, (ra, rb) in enumerate(zip(rows_a, rows_b, strict=False), 1):
                if len(ra) != len(rb):
                    diffs.append(f'{name}!R{i}: {len(ra)} colonne vs {len(rb)}')
                    continue
                for j, (va, vb) in enumerate(zip(ra, rb, strict=False), 1):
                    # L'unica cella che DEVE differire: l'ora di generazione.
                    if isinstance(va, str) and va.startswith('Generato:'):
                        skipped += 1
                        continue
                    cells += 1
                    if va != vb:
                        diffs.append(f'{name}!R{i}C{j}: {va!r} vs {vb!r}')
            lines.append(f'  foglio {name:28} {len(rows_a):>5} righe')
        lines.append(f'  celle confrontate: {cells:,} '
                     f'({skipped} saltate: timestamp di generazione)')
        if diffs:
            lines.append(f'  ✗ {len(diffs)} differenze:')
            lines.extend(f'      {d}' for d in diffs[:20])
            return False, lines
        lines.append('  ✓ allegato identico cella per cella')
        return True, lines


def main(argv=None):
    argv = list(argv if argv is not None else sys.argv[1:])
    if '--fingerprint' in argv:
        argv.remove('--fingerprint')
        if len(argv) != 1:
            print('--fingerprint vuole un solo database')
            return 2
        if not os.path.exists(argv[0]):
            print(f'✗ database non trovato: {argv[0]}')
            return 2
        for line in fingerprint(argv[0]):
            print(line)
        return 0
    year = None
    if '--annex' in argv:
        i = argv.index('--annex')
        if i + 1 >= len(argv):
            print('--annex richiede un anno')
            return 2
        year = argv[i + 1]
        del argv[i:i + 2]
    if len(argv) != 2 or argv[0] in ('-h', '--help'):
        print(__doc__)
        return 2
    ref_db, new_db = argv
    for p in (ref_db, new_db):
        if not os.path.exists(p):
            print(f'✗ database non trovato: {p}')
            return 2

    print('=' * 70)
    print('VERIFICA: il ricostruito dice le stesse cifre del riferimento?')
    print('=' * 70)
    print(f'  riferimento  {ref_db}')
    print(f'  ricostruito  {new_db}\n')

    report = compare(ref_db, new_db)
    for c in report.checks:
        mark = {'identico': '✓', 'spiegato': '~', 'DIVERSO': '✗'}[c.status]
        print(f'  {mark} {c.name:48} {c.status}')
        if c.status != 'identico':
            print(f'      {c.detail}')

    print('\n  Cifre fiscali per anno (riferimento):')
    print(f'    {"anno":6} {"match":>7} {"esenti":>7} {"plus esente":>18} {"tassabile":>14}')
    for row in report.yearly:
        print(f'    {row["year"]:6} {row["matches"]:>7} {row["exempt_n"]:>7} '
              f'{row["exempt_gain"]:>18,.2f} {row["taxable_gain"]:>14,.2f}')

    if report.notes:
        print('\n  Differenze note e spiegate:')
        for n in report.notes:
            print(f'    · {n}')

    annex_ok = True
    if year:
        print(f'\n  Allegato IRS {year}: generato da entrambi i database e '
              'confrontato')
        annex_ok, lines = compare_annex(ref_db, new_db, year)
        for ln in lines:
            print(ln)

    print('\n' + '=' * 70)
    if report.ok and annex_ok:
        print('✓ IDENTICI su tutto ciò che è confrontabile')
        print('  Le differenze marcate «spiegato» sono quelle dichiarate nel '
              'docstring di\n  questo file e in doc/rebuild.md; nessun\'altra.')
        return 0
    print('✗ DIFFERENZE NON SPIEGATE — non usare queste cifre finché non sono '
          'capite')
    return 1


if __name__ == '__main__':
    sys.exit(main())
